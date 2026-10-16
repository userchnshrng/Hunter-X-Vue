import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import CustomFieldDefinition
from app.services.imports import parse_date
from tests.template_workbook import HEADERS, RESUME_URL, workbook_bytes
from tests.test_stage_b import commit, preview, rows_of, upload

MAPPING = dict(
    enumerate(
        [
            "contacted_on",
            "name",
            "current_company_name",
            "current_title",
            "contact:phone",
            "communication",
            "custom:resume_link",
        ]
    )
)
MAPPING = {str(key): value for key, value in MAPPING.items()}


def test_template_is_blank_xlsx_and_requires_auth(context, client):
    response = client.get("/api/v1/imports/template")
    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    assert "hunter-talent-template.xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["候选人"]
    sheet = workbook.active
    assert [cell.value for cell in sheet[2]] == HEADERS
    assert all(cell.value is None for cell in sheet[1])
    assert all(cell.value is None for row in sheet.iter_rows(min_row=3) for cell in row)
    assert all(cell.comment is None and cell.hyperlink is None for row in sheet for cell in row)
    assert sheet["E3"].number_format == "@"
    assert sheet["A3"].number_format == "yyyy.mm.dd"
    workbook.close()
    uploaded = client.post(
        "/api/v1/imports/upload",
        files={
            "file": ("synthetic-blank-template.xlsx", response.content),
        },
    ).json()
    assert uploaded["sheets"][0]["headers"][1]["cells"] == HEADERS
    empty_preview = client.post(
        f"/api/v1/imports/{uploaded['id']}/preview",
        json={
            "revision": uploaded["revision"],
            "sheet_name": "候选人",
            "header_row": 2,
            "mapping": MAPPING,
        },
    )
    assert empty_preview.status_code == 422
    assert "待导入行必须" in empty_preview.json()["message"]
    client.headers.pop("Authorization")
    assert client.get("/api/v1/imports/template").status_code == 401


def test_matching_layout_preserves_hyperlink_dates_and_history(context, client):
    response = client.post(
        "/api/v1/imports/upload", files={"file": ("synthetic-template.xlsx", workbook_bytes())}
    )
    assert response.status_code == 201
    batch = preview(client, response.json(), MAPPING, header=2)
    rows = rows_of(client, batch)
    assert [row["row_number"] for row in rows] == [3, 4]
    assert all(not row["errors"] and not row["warnings"] for row in rows)
    candidate = rows[0]["normalized"]["candidate"]
    assert candidate["contacted_on"] == "2026-05-22"
    assert candidate["contact_methods"] == [{"type": "phone", "value": "00123456789"}]
    assert candidate["custom_fields"]["resume_link"] == RESUME_URL
    assert rows[0]["original"][6]["value"] == "查看简历"
    assert rows[0]["original"][6]["hyperlink"] == RESUME_URL
    assert rows[1]["normalized"]["candidate"]["contacted_on"] is None
    # Saving unrelated corrections must not replace a hyperlink with its display label.
    response = client.patch(
        f"/api/v1/imports/{batch['id']}/rows/{rows[0]['id']}",
        json={
            "revision": batch["revision"],
            "corrections": {"5": "合成修正历史沟通"},
        },
    )
    assert response.status_code == 200
    batch = response.json()
    response = commit(client, batch)
    assert response.status_code == 200
    result = response.json()
    assert result["created"] == 2
    detail = client.get(f"/api/v1/candidates/{result['candidate_ids'][0]}").json()
    assert detail["custom_fields"]["resume_link"] == RESUME_URL
    communications = detail["communications"]
    assert communications[0]["content"] == "合成修正历史沟通"
    assert communications[0]["occurred_at"] is None
    with Session(context[1].state.engine) as db:
        assert db.scalar(select(func.count()).select_from(CustomFieldDefinition)) == 0


def test_builtin_link_is_editable_value_but_fixed_definition(client):
    field = next(
        x
        for x in client.get("/api/v1/custom-fields", params={"entity_type": "candidate"}).json()
        if x["key"] == "resume_link"
    )
    assert field["builtin"] and field["enabled"] and not field["required"]
    response = client.patch(
        f"/api/v1/custom-fields/{field['id']}",
        json={
            "entity_type": "candidate",
            "label": field["label"],
            "value_type": "text",
            "enabled": False,
            "version": field["version"],
        },
    )
    assert response.status_code == 422
    assert "固定字段" in response.json()["message"]
    assert (
        client.post(
            "/api/v1/candidates",
            json={
                "name": "合成超长链接",
                "custom_fields": {"resume_link": "x" * 2001},
            },
        ).status_code
        == 422
    )


def test_matching_headers_keep_numeric_phone_and_missing_name_checks(client):
    batch = upload(client, [[], HEADERS, ["2026.05.22", "", "", "", 123456, "", ""]], xlsx=True)
    batch = preview(client, batch, MAPPING, header=2)
    row = rows_of(client, batch)[0]
    assert row["errors"] and row["warnings"]
    assert commit(client, batch).status_code == 409


@pytest.mark.parametrize("value", ["2026.02.30", "05.22", "03/04/2026"])
def test_ambiguous_or_invalid_dates_remain_errors(value):
    with pytest.raises(ValueError):
        parse_date(value)
