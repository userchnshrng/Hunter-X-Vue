"""Generate a synthetic workbook for browser tests; never read real resumes."""

import io
import sys
from datetime import date

from openpyxl import Workbook


def main():
    workbook = Workbook()
    workbook.active.title = "说明"
    workbook.active.append(["合成导入验证，请选择人才数据工作表"])
    sheet = workbook.create_sheet("人才数据")
    sheet.append(["合成工作簿 · 仅用于自动化测试"])
    sheet.merge_cells("A1:D1")
    sheet.append(
        [
            "建联日期",
            "姓名",
            "公司",
            "职位",
            "联系方式",
            "联系方式",
            "沟通记录",
            "简历来源",
            "期望地点",
            "期望年薪",
            "到岗日期",
            "意向状态",
        ]
    )
    sheet.append(
        [
            date(2026, 10, 14),
            "合成导入 · 许知远",
            "合成导入公司",
            "产品经理",
            "00123456",
            "wx-synthetic",
            "导入的历史沟通，实际沟通时间未知",
            "合成Excel",
            "杭州",
            50,
            date(2026, 12, 1),
            "可沟通",
        ]
    )
    sheet.append(
        ["03/04/2026", "=1+1", "合成导入公司", "架构师", 12345, "", "", "合成Excel", "上海", 60, "", "可沟通"]
    )
    sheet.append(
        ["", "林知远（合成人才）", "合成重复公司", "产品负责人", "", "", "", "合成Excel", "", "", "", ""]
    )
    output = io.BytesIO()
    workbook.save(output)
    sys.stdout.buffer.write(output.getvalue())


if __name__ == "__main__":
    main()
