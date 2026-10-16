import csv
import hashlib
import io
import json
import math
import posixpath
import re
import zipfile
from datetime import date, datetime

from defusedxml.ElementTree import fromstring, iterparse
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from pydantic import ValidationError
from sqlalchemy import select, text

from app.models import Candidate, Communication, ImportBatch, ImportRow, now
from app.schemas import CandidateInput
from app.services.custom_fields import definitions, snapshot, validate_fields
from app.services.records import get_record, serialize

FILE_LIMIT = 10 * 1024 * 1024
BASE_TARGETS = {
    "name",
    "contacted_on",
    "current_company_name",
    "current_title",
    "resume_source",
    "communication",
}
CONTACT_TARGETS = {f"contact:{kind}" for kind in ("phone", "wechat", "email", "other")}
REL_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def fail(message, status=422):
    raise HTTPException(status, message)


def import_root(request):
    # Tests derive this from their isolated upload root, never from the developer database directory.
    return request.app.state.settings.upload_root.parent / "imports"


def sheet_paths(archive):
    relations = {
        node.attrib["Id"]: posixpath.normpath(posixpath.join("xl", node.attrib["Target"].lstrip("/")))
        if not node.attrib["Target"].startswith("/") else node.attrib["Target"].lstrip("/")
        for node in fromstring(archive.read("xl/_rels/workbook.xml.rels"), forbid_dtd=True)
        if node.attrib.get("TargetMode") != "External"
    }
    return {
        node.attrib["name"]: relations.get(node.attrib.get(REL_ID))
        for node in fromstring(archive.read("xl/workbook.xml"), forbid_dtd=True).iter()
        if node.tag.rsplit("}", 1)[-1] == "sheet"
    }


def decode_file(data, file_format, selected=None):
    """Bound file/container/cell sizes before interpreting workbook contents."""
    if not data or len(data) > FILE_LIMIT:
        fail("导入文件必须非空且不超过 10 MB", 413)
    workbook = None
    logical_characters = 0
    try:
        if file_format == "csv":
            source = data.decode("utf-8-sig")
            if "\x00" in source:
                fail("CSV 包含非法字符，请另存为 UTF-8 CSV")
            csv.field_size_limit(20000)
            rows = []
            for number, row in enumerate(csv.reader(io.StringIO(source)), 1):
                if number > 10000 or len(row) > 100:
                    fail("文件超过 10,000 行或 100 列，请拆分后上传")
                rows.append((number, [{"value": v, "type": "f" if v.startswith("=") else "s"} for v in row]))
            tables = {"CSV": rows}
        else:
            hyperlinks = {}
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(x.file_size for x in entries) > 50 * 1024 * 1024:
                    fail("Excel 展开后超过 50 MB 或文件条目过多")
                if any("vbaproject" in x.filename.lower() for x in entries):
                    fail("不支持含宏的工作簿，请另存为普通 .xlsx")
                if "xl/workbook.xml" not in archive.namelist():
                    fail("文件不是有效的 Excel 工作簿")
                paths = sheet_paths(archive)
                # Reject sparse or dishonest dimension declarations too; never silently truncate cells.
                for entry in entries:
                    if entry.filename in paths.values():
                        rel_path = posixpath.join(
                            posixpath.dirname(entry.filename), "_rels",
                            posixpath.basename(entry.filename) + ".rels",
                        )
                        relations = {}
                        if rel_path in archive.namelist():
                            relations = {
                                node.attrib["Id"]: node.attrib.get("Target", "")
                                for node in fromstring(archive.read(rel_path), forbid_dtd=True)
                                if node.attrib.get("Type", "").endswith("/hyperlink")
                                and node.attrib.get("TargetMode") == "External"
                            }
                        links = hyperlinks[entry.filename] = {}
                        with archive.open(entry) as stream:
                            for _, node in iterparse(stream, events=("end",), forbid_dtd=True):
                                tag = node.tag.rsplit("}", 1)[-1]
                                if tag == "row" and int(node.attrib.get("r", 0)) > 10000:
                                    fail("工作表超过 10,000 行，请拆分后上传")
                                if tag == "c":
                                    coordinate = re.fullmatch(r"([A-Z]+)(\d+)", node.attrib.get("r", ""))
                                    if coordinate and (
                                        column_index_from_string(coordinate[1]) > 100
                                        or int(coordinate[2]) > 10000
                                    ):
                                        fail("工作表超过 10,000 行或 100 列")
                                if tag == "hyperlink":
                                    target = relations.get(node.attrib.get(REL_ID), "")
                                    if len(target) > 20000:
                                        fail("超链接超过 20,000 字")
                                    if target:
                                        # Keep targets as text; never resolve or fetch external resources.
                                        links[node.attrib.get("ref", "")] = target
                                node.clear()
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            if len(workbook.sheetnames) > 50:
                fail("工作表超过 50 张，请拆分后上传")
            tables = {}
            for sheet in workbook:
                if selected is not None and sheet.title != selected:
                    continue
                if len(sheet.title) > 100:
                    fail("工作表名称过长")
                sheet.reset_dimensions()
                rows = []
                max_row = 10000 if selected else 20
                for number, row in enumerate(sheet.iter_rows(max_row=max_row, max_col=100), 1):
                    cells = []
                    for cell in row:
                        value = cell.value
                        if isinstance(value, (date, datetime)):
                            value = (
                                value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                            )
                        value = "" if value is None else str(value)
                        logical_characters += len(value)
                        if logical_characters > 10 * 1024 * 1024:
                            fail("工作簿展开文字过多，请拆分后上传")
                        if len(value) > 20000:
                            fail(f"第 {number} 行单元格超过 20,000 字")
                        cells.append(
                            {"value": value, "type": cell.data_type if cell.value is not None else "s"}
                        )
                        coordinate = getattr(cell, "coordinate", None)
                        link = hyperlinks.get(paths.get(sheet.title), {}).get(coordinate)
                        if link:
                            cells[-1]["hyperlink"] = link
                            logical_characters += len(link)
                            if logical_characters > 10 * 1024 * 1024:
                                fail("工作簿展开文字过多，请拆分后上传")
                    while cells and cells[-1]["value"] == "" and not cells[-1].get("hyperlink"):
                        cells.pop()
                    rows.append((number, cells))
                tables[sheet.title] = rows
        if selected is not None:
            if selected not in tables:
                fail("工作表不存在，请重新选择")
            return tables[selected]
        return [
            {
                "name": name,
                "headers": [{"row": n, "cells": [x["value"] for x in cells]} for n, cells in rows[:20]],
            }
            for name, rows in tables.items()
        ]
    except HTTPException:
        raise
    except (UnicodeError, csv.Error):
        fail("CSV 无法读取，请使用 UTF-8 编码并检查单元格长度")
    except Exception:
        # Parser exceptions may contain workbook text or filesystem details.
        fail("Excel 无法读取：文件损坏、加密、格式不符或包含不安全的 XML")
    finally:
        if workbook:
            workbook.close()


def write_transaction(db):
    # Authentication has already opened a read transaction. Start a fresh write snapshot.
    db.rollback()
    if db.bind.dialect.name == "sqlite":
        db.execute(text("BEGIN IMMEDIATE"))


def owned_batch(db, record_id, user_id, mutable=False):
    batch = (
        db.scalar(select(ImportBatch).where(ImportBatch.id == record_id).with_for_update())
        if mutable
        else db.get(ImportBatch, record_id)
    )
    if not batch or batch.created_by != user_id:
        fail("导入批次不存在", 404)
    return batch


def ensure_editable(batch, revision=None):
    if batch.status in ("committed", "cancelled"):
        fail("该批次已提交或取消，不能再次修改", 409)
    if revision is not None and revision != batch.revision:
        fail("预览已更新，请刷新批次后重试", 409)


def active_rows(db, batch):
    return db.scalars(
        select(ImportRow)
        .where(
            ImportRow.batch_id == batch.id,
            ImportRow.sheet_name == batch.sheet_name,
            ImportRow.row_number > batch.header_row,
        )
        .order_by(ImportRow.row_number)
    ).all()


def parse_date(value):
    value = value.strip()
    if not value:
        return None
    match = re.fullmatch(r"(\d{4})(?:-|/|\.|年)(\d{1,2})(?:-|/|\.|月)(\d{1,2})日?", value)
    if not match:
        raise ValueError("日期应为明确的年、月、日，例如 2026-10-16")
    return date(*map(int, match.groups())).isoformat()


def normalize_row(db, row, mapping, fields):
    values = {"name": "", "contact_methods": [], "custom_fields": {}}
    errors, warnings = [], []
    communication = ""
    for column, target in mapping.items():
        if not target:
            continue
        index = int(column)
        cell = row.original[index] if index < len(row.original) else {"value": "", "type": "s"}
        original = cell.get("hyperlink", cell["value"]) if target == "custom:resume_link" else cell["value"]
        value = row.corrections.get(column, original).strip()
        try:
            if column not in row.corrections and cell["type"] in ("f", "e"):
                raise ValueError("公式或错误单元格，请转换为普通值或在预览中修正")
            if target.startswith("contact:"):
                if value:
                    values["contact_methods"].append({"type": target.split(":")[1], "value": value})
                    if column not in row.corrections and (
                        cell["type"] == "n" or re.fullmatch(r"[+-]?\d+(\.\d+)?[Ee][+-]?\d+", value)
                    ):
                        warnings.append(
                            {
                                "column": column,
                                "message": "数字联系方式需核对科学计数法及前导零；丢失内容无法恢复",
                            }
                        )
            elif target.startswith("custom:"):
                key = target.split(":", 1)[1]
                field = fields[key]
                if field.value_type == "date":
                    value = parse_date(value)
                elif field.value_type == "number":
                    value = float(value) if value else None
                    if value is not None and not math.isfinite(value):
                        raise ValueError("数字必须为有限数值")
                elif field.value_type == "select" and value:
                    matches = [
                        x["value"] for x in field.options if x["label"] == value or x["value"] == value
                    ]
                    if len(set(matches)) != 1:
                        raise ValueError("单选值无效或显示名不唯一，请使用明确选项")
                    value = matches[0]
                values["custom_fields"][key] = value
            elif target == "communication":
                if len(value) > 20000:
                    raise ValueError("沟通记录不能超过 20,000 字")
                communication = value
            else:
                values[target] = parse_date(value) if target == "contacted_on" else value
        except (ValueError, KeyError) as error:
            errors.append(
                {
                    "column": column,
                    "message": str(error) if not isinstance(error, KeyError) else "字段定义已变更",
                }
            )
    try:
        body = CandidateInput.model_validate(values)
        values = body.model_dump(mode="json")
        values["custom_fields"] = validate_fields(
            db, "candidate", values["custom_fields"], field_definitions=list(fields.values())
        )
    except ValidationError as error:
        errors.extend(
            {"column": str(x["loc"][0]), "message": "必填内容缺失、过长或格式不正确"} for x in error.errors()
        )
    except HTTPException as error:
        errors.append({"column": "custom_fields", "message": str(error.detail)})
    row.normalized = {"candidate": values, "communication": communication}
    row.errors, row.warnings = errors, warnings


def contact_key(contact):
    value = contact["value"].strip()
    if contact["type"] == "phone":
        value = re.sub(r"[\s()-]", "", value)
    elif contact["type"] == "email":
        value = value.casefold()
    return value


def duplicate_index(db, rows):
    names, contacts = {}, {}
    for item in db.scalars(select(Candidate)):
        reference = {"kind": "existing", "id": item.id, "name": item.name}
        names.setdefault(item.name.casefold(), []).append(reference)
        for contact in item.contact_methods:
            contacts.setdefault(contact_key(contact), []).append(reference)
    for row in rows:
        value = row.normalized["candidate"]
        reference = {"kind": "batch", "row": row.row_number}
        names.setdefault(value.get("name", "").casefold(), []).append(reference)
        for contact in value.get("contact_methods", []):
            contacts.setdefault(contact_key(contact), []).append(reference)
    return names, contacts


def revalidate(db, batch):
    rows = active_rows(db, batch)
    fields = {x.key: x for x in definitions(db, "candidate") if x.enabled}
    for row in rows:
        normalize_row(db, row, batch.mapping, fields)
    names, contacts = duplicate_index(db, rows)
    counts = {"total": len(rows), "errors": 0, "review": 0, "skipped": 0, "ready": 0}
    for row in rows:
        value = row.normalized["candidate"]
        matches = list(names.get(value.get("name", "").casefold(), []))
        for contact in value.get("contact_methods", []):
            matches += contacts.get(contact_key(contact), [])
        matches = [x for x in matches if not (x["kind"] == "batch" and x["row"] == row.row_number)]
        duplicates = sorted({json.dumps(x, sort_keys=True) for x in matches})
        duplicates = [json.loads(x) for x in duplicates]
        if duplicates != row.duplicates:
            row.decision = ""
        row.duplicates = duplicates
        if row.errors:
            row.status = "error"
            counts["errors"] += 1
        elif duplicates and row.decision == "skip":
            row.status = "skip"
            counts["skipped"] += 1
        elif (duplicates and row.decision != "create") or (row.warnings and not row.acknowledged):
            row.status = "review"
            counts["review"] += 1
        else:
            row.status = "ready"
            counts["ready"] += 1
    batch.counts = counts
    batch.status = "ready" if counts["total"] and not counts["errors"] and not counts["review"] else "review"
    return rows


def preview(db, batch, body, root):
    ensure_editable(batch, body.revision)
    allowed = (
        BASE_TARGETS
        | CONTACT_TARGETS
        | {f"custom:{x.key}" for x in definitions(db, "candidate") if x.enabled}
    )
    seen = set()
    for column, target in body.mapping.items():
        if not re.fullmatch(r"\d{1,2}", column) or int(column) >= 100 or str(int(column)) != column:
            fail("列编号应在 0–99 之间")
        if target and (target not in allowed or (target in seen and target not in CONTACT_TARGETS)):
            fail("映射字段不存在或被重复映射")
        seen.add(target)
    if "name" not in seen:
        fail("必须映射姓名列")
    records = decode_file((root / batch.storage_key).read_bytes(), batch.file_format, body.sheet_name)
    records = [
        (number, cells)
        for number, cells in records
        if number > body.header_row and any(x["value"].strip() for x in cells)
    ]
    if not records or len(records) > 5000:
        fail("待导入行必须为 1–5,000 条，请检查表头或拆分文件")
    changed = (
        batch.mapping != body.mapping
        or batch.sheet_name != body.sheet_name
        or batch.header_row != body.header_row
    )
    existing = {
        (x.sheet_name, x.row_number): x
        for x in db.scalars(select(ImportRow).where(ImportRow.batch_id == batch.id))
    }
    for number, cells in records:
        row = existing.get((body.sheet_name, number))
        if row is None:
            row = ImportRow(batch_id=batch.id, sheet_name=body.sheet_name, row_number=number, original=cells)
            db.add(row)
        if changed:
            row.corrections, row.decision, row.acknowledged = {}, "", False
    batch.sheet_name, batch.header_row, batch.mapping = body.sheet_name, body.header_row, body.mapping
    batch.field_snapshot = snapshot(db)
    batch.revision += 1
    db.flush()
    revalidate(db, batch)
    db.commit()
    return serialize(batch)


def edit_row(db, batch, row_id, body):
    ensure_editable(batch, body.revision)
    row = get_record(db, ImportRow, row_id)
    if row.batch_id != batch.id or row.sheet_name != batch.sheet_name or row.row_number <= batch.header_row:
        fail("预览行不存在", 404)
    if batch.field_snapshot != snapshot(db):
        fail("字段定义已更新，请重新生成预览", 409)
    if body.corrections is not None:
        if any(
            column not in batch.mapping or len(value) > 20000 for column, value in body.corrections.items()
        ):
            fail("只能修正映射列，单元格最多 20,000 字")
        row.corrections = {**row.corrections, **body.corrections}
        row.decision, row.acknowledged = "", False
    # Refresh duplicates before applying the user's decision, so new duplicates cannot reuse stale approval.
    revalidate(db, batch)
    if body.decision is not None:
        if body.decision == "skip" and not row.duplicates:
            fail("只能显式跳过重复行；错误行需要先修正")
        row.decision = body.decision
    if body.acknowledged is not None:
        row.acknowledged = body.acknowledged
    revalidate(db, batch)
    batch.revision += 1
    db.commit()
    return serialize(batch)


def commit(db, batch, body):
    digest = hashlib.sha256(f"{batch.id}:{body.revision}".encode()).hexdigest()
    reused = db.scalar(
        select(ImportBatch).where(
            ImportBatch.created_by == batch.created_by, ImportBatch.idempotency_key == body.idempotency_key
        )
    )
    if reused and (reused.id != batch.id or reused.commit_digest != digest):
        fail("提交标识已用于其他确认内容", 409)
    if batch.status == "committed":
        if batch.commit_digest != digest:
            fail("该批次已按其他版本提交", 409)
        return batch.result
    ensure_editable(batch, body.revision)
    if batch.field_snapshot != snapshot(db):
        fail("字段定义已更新，请重新生成预览", 409)
    rows = revalidate(db, batch)
    if batch.status != "ready":
        # Persist newly detected duplicates/errors, not partial business records.
        batch.revision += 1
        db.commit()
        fail("请先修正错误、核对提醒并处理重复行；预览已更新", 409)
    result = {"batch_id": batch.id, "created": 0, "skipped": 0, "candidate_ids": []}
    for row in rows:
        if row.status == "skip":
            result["skipped"] += 1
            continue
        candidate = Candidate(**CandidateInput.model_validate(row.normalized["candidate"]).model_dump())
        db.add(candidate)
        db.flush()
        if row.normalized["communication"]:
            db.add(
                Communication(
                    candidate_id=candidate.id,
                    content=row.normalized["communication"],
                    channel="导入记录",
                    occurred_at=None,
                )
            )
        row.candidate_id, row.status = candidate.id, "created"
        result["candidate_ids"].append(candidate.id)
        result["created"] += 1
    batch.result, batch.status = result, "committed"
    batch.idempotency_key, batch.commit_digest = body.idempotency_key, digest
    batch.committed_at = now()
    db.commit()
    return result


def safe_csv(value):
    value = str(value)
    return "'" + value if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) else value
